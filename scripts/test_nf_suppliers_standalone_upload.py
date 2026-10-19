"""Teste de regressão da importação isolada de fornecedores na Contingência Manual.

Usa apenas arquivo XLTX sintético e inspeção AST; não altera base real nem acessa API.
"""
import ast
import io
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook


source = Path("streamlit_app.py").read_text(encoding="utf-8")
tree = ast.parse(source)
by_name = {
    node.name: node
    for node in tree.body
    if isinstance(node, ast.FunctionDef)
}

feed = by_name["render_suppliers_feed"]
assert [arg.arg for arg in feed.args.args] == ["upload"]
assert len(feed.args.defaults) == 1
assert isinstance(feed.args.defaults[0], ast.Constant)
assert feed.args.defaults[0].value is None

process = by_name["render_file_processing"]
standalone = [
    node for node in ast.walk(process)
    if isinstance(node, ast.Call)
    and isinstance(node.func, ast.Name)
    and node.func.id == "render_suppliers_feed"
]
assert len(standalone) == 1, "A página precisa expor a validação independente."
assert any(
    kw.arg == "upload"
    and isinstance(kw.value, ast.Name)
    and kw.value.id == "supplier_file"
    for kw in standalone[0].keywords
), "A validação deve reutilizar o arquivo FORNECEDORES já selecionado."

process_source = ast.get_source_segment(source, process)
feed_source = ast.get_source_segment(source, feed)
assert process_source.index("render_suppliers_feed(upload=supplier_file)") < process_source.index("if analyze_bases:")
assert "if supplier_file is not None:" in process_source
assert "supplier_import_fingerprint" in feed_source
assert "hashlib.sha256(supplier_bytes).hexdigest()" in feed_source
assert '"VALIDAR BASE DE FORNECEDORES"' in feed_source
assert '"SUBSTITUIR BASE DE FORNECEDORES"' in feed_source
assert "disabled=not can_replace" in feed_source
assert "db.replace_suppliers(" in feed_source

# Isola a função de leitura, sem iniciar a aplicação Streamlit.
reader = by_name["_read_uploaded_table_cached"]
reader.decorator_list = []
module = ast.Module(body=[reader], type_ignores=[])
ast.fix_missing_locations(module)
namespace = {
    "pd": pd,
    "Path": Path,
    "io": io,
    "load_workbook": load_workbook,
}
exec(compile(module, "<supplier-xltx-test>", "exec"), namespace)
read = namespace["_read_uploaded_table_cached"]

book = Workbook()
book.template = True
sheet = book.active
sheet.title = "FORNECEDORES"
sheet.append([f"COL{i}" for i in range(1, 16)])
data = [""] * 15
data[0] = "001234"
data[1] = "01"
data[2] = "FORNECEDOR TESTE LTDA"
data[3] = "TESTE"
data[10] = "J"
data[14] = "12.345.678/0001-90"
sheet.append(data)
buffer = io.BytesIO()
book.save(buffer)
raw = buffer.getvalue()
frame, sheets = read(raw, "fornecedores.xltx", None, None)
assert sheets == ["FORNECEDORES"]
assert frame.shape == (2, 15), frame.shape
assert frame.iloc[1, 0] == "001234"
assert frame.iloc[1, 14] == "12.345.678/0001-90"

try:
    read(b"arquivo-corrompido", "corrompido.xltx", None, None)
except ValueError as exc:
    assert "Não foi possível abrir o arquivo .xltx" in str(exc)
else:
    raise AssertionError("Um arquivo XLTX inválido não pode ser aceito.")

print("NFS_SUPPLIERS_STANDALONE_UPLOAD_OK")
