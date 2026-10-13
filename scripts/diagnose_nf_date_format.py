"""Diagnóstico controlado do layout STSUP01 na Central SETTA.

Não imprime fornecedor, NF, produto, CNPJ nem valores de células.
Exibe apenas contagens, tipos de dados e posição provável da coluna de data.
"""
from __future__ import annotations

from collections import Counter
from datetime import datetime, date
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from central_nfs_data import download_normalized_source, download_source_bytes, source_raw_frame


def classify(value):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "VAZIA"
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return "DATA_EXCEL"
    if isinstance(value, (int, float)):
        number = float(value)
        if 20000 <= number < 100000:
            return "SERIAL_EXCEL"
        if number > 1e9:
            return "EPOCH_POSSIVEL"
        return "NUMERICO"
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return "VAZIA"
        if pd.notna(pd.to_datetime(s, dayfirst=True, errors="coerce")):
            return "DATA_TEXTO"
        return "TEXTO"
    return type(value).__name__


def report(label, sample):
    print("FONTE", label, "LINHAS_ANALISADAS", len(sample))
    print("TOTAL_COLUNAS", max((len(x) for x in sample), default=0))
    for index in [0, 1, 2, 3, 4, 5, 9, 12, 30]:
        if index >= max((len(x) for x in sample), default=0):
            print("COLUNA", index + 1, "AUSENTE")
            continue
        types = Counter(classify(row[index] if len(row) > index else None) for row in sample)
        print("COLUNA", index+1, "TIPOS", dict(types))


try:
    pack, meta = download_normalized_source("nf")
    frame = source_raw_frame(pack, 0)
    print("NORMALIZADA", "LINHAS", len(frame), "COLUNAS", frame.shape[1])
    report("NORMALIZADA", [tuple(row) for row in frame.iloc[2:502].itertuples(index=False, name=None)])
except Exception as exc:
    print("NORMALIZADA_ERRO", type(exc).__name__, str(exc)[:150])

try:
    raw, meta = download_source_bytes("nf")
    name = str(meta.get("last_file_name") or "nf.xlsx")
    print("ORIGINAL", "EXTENSAO", name.rsplit(".",1)[-1].lower(), "BYTES", len(raw))
    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        print("ORIGINAL", "ABA", ws.title, "TOTAL_COLUNAS", ws.max_column, "TOTAL_LINHAS", ws.max_row)
        values = []
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if idx <= 2: continue
            values.append(tuple(row))
            if len(values)>=500: break
        report("ORIGINAL", values)
    finally:
        wb.close()
except Exception as exc:
    print("ORIGINAL_ERRO", type(exc).__name__, str(exc)[:150])
