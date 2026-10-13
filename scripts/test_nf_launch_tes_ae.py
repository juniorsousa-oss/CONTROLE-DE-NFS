"""Regressão: STSUP01 só confirma NF-e com TES AE em todos os itens."""
from __future__ import annotations

import ast
from datetime import datetime, timedelta, timezone
import io
from pathlib import Path
import re
from types import SimpleNamespace

import pandas as pd
from openpyxl import Workbook, load_workbook
from nf_launch_validation import tes_preenchida, validar_lancamento_nf

ROOT = Path(__file__).resolve().parents[1]
source = (ROOT / "streamlit_app.py").read_text(encoding="utf-8")
tree = ast.parse(source)


def normalize(value):
    return " ".join(str(value or "").strip().upper().split())


def code(value):
    digits = "".join(x for x in str(value or "") if x.isdigit())
    return digits.lstrip("0") or ("0" if digits else "")


def digits(value):
    return "".join(x for x in str(value or "") if x.isdigit())


def similarity(a, b):
    # Aproximações são apenas marcadas para revisão.
    return 95 if normalize(a).startswith("ALFA") and normalize(b).startswith("ALFA") else 0


def verify(record, rows):
    return validar_lancamento_nf(
        record, rows, normalizar_nome=normalize, similaridade=similarity,
        normalizar_codigo=code, somente_digitos=digits,
    )


base = {"numero_nf": "123", "fornecedor_padrao": "ALFA COMPONENTES", "cnpj_fornecedor": "11111111000191"}
def row(tes, **kwargs):
    return {"numero_nf": "123", "fornecedor": "ALFA COMPONENTES",
            "fornecedor_codigo": "00031", "cnpj_resolvido": "11111111000191", "tes": tes, **kwargs}

for empty in (None, "", "    ", float("nan"), "None", "NaN", "null", 0, "0"):
    assert not tes_preenchida(empty), empty
    assert verify(base, [row(empty)]) == "AGUARDANDO_LANCAMENTO"

assert tes_preenchida("501")
assert tes_preenchida(501)
assert verify(base, [row("501"), row("503")]) == "CONFIRMADO"
assert verify(base, [row("501"), row(None)]) == "AGUARDANDO_LANCAMENTO"
assert verify(base, []) == "NAO_ENCONTRADA"
assert verify(base, [row("501", **{})]) == "CONFIRMADO"

diff_cnpj = [{**row("501"), "cnpj_resolvido": "99999999000199"}]
assert verify(base, diff_cnpj) == "CONFERIR_FORNECEDOR"

fuzzy = [{**row("501"), "fornecedor": "ALFA COMPONENTES ELETRICOS", "cnpj_resolvido": ""}]
assert verify(base, fuzzy) == "CONFERIR_FORNECEDOR"

blank_name = {**base, "fornecedor_padrao": "", "cnpj_fornecedor": ""}
assert verify(blank_name, [row("501")]) == "CONFERIR_FORNECEDOR"

# Outro fornecedor, mesmo número, TES presente não confirma NF da ALFA.
different = [{"numero_nf": "123", "fornecedor": "BETA", "fornecedor_codigo": "32",
             "cnpj_resolvido": "22222222000199", "tes": "501"}]
assert verify(base, different) == "NAO_ENCONTRADA"

# Quando o CNPJ está confirmado, nomes diferentes pertencem à mesma NF:
# qualquer item sem TES bloqueia o conjunto.
mixed = [row("501"), {**row(""), "fornecedor": "ALFA COMPONENTES LTDA"}]
assert verify(base, mixed) == "AGUARDANDO_LANCAMENTO"


def extract_ast(name, scope):
    matches = [x for x in tree.body if isinstance(x, ast.FunctionDef) and x.name == name]
    assert len(matches) == 1, name
    function = matches[0]
    function.decorator_list = []
    exec(compile(ast.Module(body=[function], type_ignores=[]), "<nfs-extract>", "exec"), scope)
    return scope[name]


def excel_frame_rows():
    wb = Workbook()
    ws = wb.active
    ws.append(["c" + str(i) for i in range(31)])
    ws.append(["x" for _ in range(31)])
    for nf, supplier, tes in ((123, "ALFA COMPONENTES", None),
                              (123, "ALFA COMPONENTES", 501),
                              (124, "BETA", 502)):
        cells = [None] * 31
        cells[3] = nf
        cells[4] = "00031"
        cells[5] = supplier
        cells[30] = tes
        ws.append(cells)
    target = io.BytesIO()
    wb.save(target)
    return target.getvalue(), [list(row_) for row_ in ws.values]


raw_xlsx, rows_excel = excel_frame_rows()
scope = {
    "pd": pd, "Path": Path, "io": io, "load_workbook": load_workbook,
    "normalized_nf": code, "_supplier_code_norm": code,
    "supplier_validation_name": normalize,
    "_read_uploaded_table_cached": lambda raw, filename, a, b:
        (pd.DataFrame(rows_excel[2:]), []),
}
scope["_stsup01_tes_cell"] = extract_ast("_stsup01_tes_cell", scope)
scope["_stsup01_operational_date"] = extract_ast("_stsup01_operational_date", scope)
read_manual = extract_ast("_extract_launch_report_cached", scope)
read_pack = extract_ast("_extract_launch_report_pack", {
    **scope, "central_data": SimpleNamespace(source_raw_frame=lambda pack, i: pd.DataFrame(pack["rows"])),
})
manual = read_manual(raw_xlsx, "STSUP01.xlsx")
normalized = read_pack({"rows": rows_excel})
for result in (manual, normalized):
    assert len(result[result["numero_nf"] == "123"]) == 2, result
    assert result.columns.tolist() == ["numero_nf", "fornecedor_codigo", "fornecedor", "tes"]
    assert not tes_preenchida(result[result["numero_nf"] == "123"]["tes"].iloc[0])
    assert tes_preenchida(result[result["numero_nf"] == "123"]["tes"].iloc[1])
    assert verify(base, [{**r, "cnpj_resolvido": "11111111000191"} for r in
                         result[result["numero_nf"] == "123"].to_dict("records")]) == "AGUARDANDO_LANCAMENTO"


# A normalização pode eliminar AE quando todos os TES estão em branco:
# esta ausência deve gerar base pendente e não falha da Central.
raw_without_ae = [line[:30] for line in rows_excel]
missing_ae_scope = {
    **scope,
    "central_data": SimpleNamespace(
        source_raw_frame=lambda pack, i: pd.DataFrame(pack["rows"]),
    ),
    "_read_uploaded_table_cached": lambda raw, filename, a, b:
        (pd.DataFrame(raw_without_ae[2:]), []),
}
read_pack_missing = extract_ast("_extract_launch_report_pack", missing_ae_scope)
read_manual_missing = extract_ast("_extract_launch_report_cached", missing_ae_scope)
pack_without_ae = read_pack_missing({"rows": raw_without_ae})
table_without_ae = read_manual_missing(b"irrelevant", "STSUP01.csv")
for frame in (pack_without_ae, table_without_ae):
    assert not frame.empty
    assert frame["tes"].eq("").all()
    rows_match = [{**row, "cnpj_resolvido": "11111111000191"}
                  for row in frame[frame["numero_nf"] == "123"].to_dict("records")]
    assert verify(base, rows_match) == "AGUARDANDO_LANCAMENTO"

# O parser de MRP deve manter NFs com TES omitida, e não esvaziar a base.
clean_pack_scope = {
    **scope,
    "date": __import__("datetime").date, "timedelta": __import__("datetime").timedelta,
    "central_data": SimpleNamespace(source_raw_frame=lambda pack, i: pd.DataFrame(pack["rows"])),
    "normalized_material_code": lambda value: str(value or "").strip(),
    "re": re,
}
clean_pack_scope["_stsup01_operational_date"] = scope["_stsup01_operational_date"]
read_clean_pack = extract_ast("_clean_mrp_nf_pack", clean_pack_scope)
# Simular uma linha operacional com dados de A a M, sem coluna AE.
matrix = [list("HEAD" for _ in range(13)), ["" for _ in range(13)]]
record = [None] * 13
record[0] = datetime(2026, 10, 13)
record[3] = 123
record[4] = "00031"
record[5] = "ALFA COMPONENTES"
record[11] = "1005"
record[12] = "PARAFUSO"
matrix.append(record)
mrp_rows, _ = read_clean_pack({"rows": matrix}, "2026-10-13")
assert len(mrp_rows) == 1 and mrp_rows.iloc[0]["tes"] == ""

# Regressão real: normalizador SheetJS usa cellDates=False e converte datas
# para número serial, diferentemente do openpyxl.
from openpyxl.utils.datetime import to_excel

excel_serial = to_excel(datetime(2026, 10, 13))
assert excel_serial > 46000
parsed_serial = scope["_stsup01_operational_date"](excel_serial)
assert parsed_serial.date().isoformat() == "2026-10-13"
assert scope["_stsup01_operational_date"](str(excel_serial)).date().isoformat() == "2026-10-13"
assert scope["_stsup01_operational_date"]("13/10/2026").date().isoformat() == "2026-10-13"
assert pd.isna(scope["_stsup01_operational_date"]("INDEFINIDA"))

serial_matrix = [list("HEADER" for _ in range(31)), ["" for _ in range(31)]]
serial_row = [None] * 31
serial_row[0] = excel_serial
serial_row[3] = 123
serial_row[4] = "00031"
serial_row[5] = "ALFA COMPONENTES"
serial_row[11] = "1005"
serial_row[12] = "PARAFUSO"
serial_row[30] = None
serial_matrix.append(serial_row)
serial_rows, serial_stats = read_clean_pack({"rows": serial_matrix}, "2026-10-13")
assert len(serial_rows) == 1, serial_stats
assert serial_rows.iloc[0]["data_pre_nota"].isoformat() == "2026-10-13"
assert serial_stats["datas_invalidas"] == 0, serial_stats
assert serial_stats["linhas_30_dias_tes"] == 1

# Data realmente inválida: impedir esvaziamento silencioso da base MRP.
serial_matrix[-1][0] = "DATA_INCORRETA"
try:
    read_clean_pack({"rows": serial_matrix}, "2026-10-13")
    raise AssertionError("Datas inválidas não podem produzir MRP com status ATUALIZADO")
except ValueError as exc:
    assert "datas inválidas" in str(exc)


# Testar reconciliação no modo de testes: uma confirmação antiga é revogada
# se o STSUP01 passa a informar TES vazio.
now = datetime.now(timezone.utc)
manifest = [{
    "id": "nf1", "numero_nf": "123", "fornecedor_padrao": "ALFA COMPONENTES",
    "cnpj_fornecedor": "11111111000191",
    "tipo_documento": "NF-e", "enviado_em": (now - timedelta(hours=2)).isoformat(),
    "lancado_em": now.isoformat(), "lancamento_verificado_em": None,
}]
class AttrState(dict):
    def __getattr__(self, name):
        return self[name]

    def __setattr__(self, name, value):
        self[name] = value


session_state = AttrState({"current_test_manifest": manifest})
st = SimpleNamespace(session_state=session_state)
reconcile_scope = {
    "pd": pd, "st": st, "db": SimpleNamespace(configured=lambda: False),
    "SAVE_NF_HISTORY": False, "is_cte_document_type": lambda x: str(x).upper().startswith("CT"),
    "normalized_nf": code, "validar_lancamento_nf": validar_lancamento_nf,
    "_supplier_cnpj_lookup": lambda df: (pd.Series(["11111111000191"] * len(df), index=df.index), 0),
    "supplier_validation_name": normalize, "supplier_similarity": similarity,
    "_supplier_code_norm": code, "digits_only": digits,
    "now_local": lambda: now,
}
reconcile = extract_ast("reconcile_launch_report", reconcile_scope)
result = reconcile(pd.DataFrame([{"numero_nf": "123", "fornecedor_codigo": "00031",
                                 "fornecedor": "ALFA COMPONENTES", "tes": ""}]))
assert result["lancados"] == 0 and result["verificados"] == 1, result
assert st.session_state["current_test_manifest"][0]["lancado_em"] is None

result = reconcile(pd.DataFrame([{"numero_nf": "123", "fornecedor_codigo": "00031",
                                 "fornecedor": "ALFA COMPONENTES", "tes": "501"}]))
assert result["lancados"] == 1, result
assert st.session_state["current_test_manifest"][0]["lancado_em"]

print("NF_TES_AE_LAUNCH_VALIDATION_OK")
